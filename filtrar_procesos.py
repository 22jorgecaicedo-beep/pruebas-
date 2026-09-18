"""
Filtra un informe de procesos concursales y deja solo los procesos que NO
estan en estado "ACUERDO CONFIRMADO".

Lee el .xlsx del informe y genera un Excel nuevo con:

  - Una hoja por cada hoja del informe (REORGANIZACION, LIQUIDACION, ...),
    con sus propios encabezados y solo las filas que quedaron, mas el
    numero de fila del informe original para poder verificar.
  - Una hoja "Resumen" con los totales por hoja, el detalle por estado y
    que columna se uso para decidir en cada caso.

Uso (Windows):

    python filtrar_procesos.py "C:\\Users\\Owner\\Downloads\\informe procesos concursales Agosto 2026.xlsx"

Si no se pasa ninguna ruta, busca en la carpeta de Descargas el informe
mas reciente cuyo nombre empiece por "informe procesos concursales".
"""

import argparse
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================= CONFIGURACION =============================

# Estado que se quiere DEJAR POR FUERA del reporte.
ESTADO_A_EXCLUIR = "ACUERDO CONFIRMADO"

# Nombres posibles de la columna de estado, EN ORDEN DE PREFERENCIA.
# Ojo: un informe puede tener varias columnas que empiecen por "ESTADO"
# (por ejemplo "ESTADO DEL PROCESO LEASING...") que no son el estado del
# proceso concursal; por eso el orden importa y ademas se revisa el
# contenido de la columna antes de elegirla.
NOMBRES_COLUMNA_ESTADO = [
    "ETAPA ACTUAL",
    "ESTADO DEL PROCESO",
    "ESTADO PROCESO",
    "ESTADO ACTUAL",
    "ESTADO DEL TRAMITE",
    "ETAPA DEL PROCESO",
    "ETAPA",
    "ESTADO",
]

# Prefijos con los que puede empezar una columna de estado, cuando su
# nombre exacto no esta en la lista de arriba.
PREFIJOS_COLUMNA_ESTADO = ("ESTADO", "ETAPA")

# Cuantas filas se revisan al inicio de la hoja buscando el encabezado.
MAX_FILAS_ENCABEZADO = 30

# ===========================================================================


def normalizar(valor) -> str:
    """Texto en mayusculas, sin tildes y con los espacios colapsados."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().upper()


def fila_vacia(valores) -> bool:
    return all(v is None or str(v).strip() == "" for v in valores)


def limpiar_celda(valor):
    """Deja el valor listo para escribirlo en el Excel de salida."""
    if isinstance(valor, (datetime, date, int, float, bool)) or valor is None:
        return valor
    return str(valor).strip()


def es_acuerdo_confirmado(estado, exacto: bool) -> bool:
    estado_normalizado = normalizar(estado)
    objetivo = normalizar(ESTADO_A_EXCLUIR)
    return estado_normalizado == objetivo if exacto else objetivo in estado_normalizado


def columnas_candidatas(encabezados_normalizados):
    """
    Devuelve [(prioridad, indice)] de las columnas que podrian ser la de
    estado. Menor prioridad = mejor candidata: primero las que coinciden
    exactamente con un nombre conocido (en el orden de la lista), y
    despues las que solo empiezan por "ESTADO" o "ETAPA".
    """
    candidatas = []
    for indice, encabezado in enumerate(encabezados_normalizados):
        if not encabezado:
            continue
        if encabezado in NOMBRES_COLUMNA_ESTADO:
            candidatas.append((NOMBRES_COLUMNA_ESTADO.index(encabezado), indice))
        elif encabezado.startswith(PREFIJOS_COLUMNA_ESTADO):
            candidatas.append((len(NOMBRES_COLUMNA_ESTADO) + indice, indice))
    return sorted(candidatas)


def elegir_columna_estado(candidatas, filas_datos, exacto: bool):
    """
    Entre las columnas candidatas, prefiere una cuyo CONTENIDO realmente
    tenga el estado buscado; si ninguna lo tiene (puede pasar en una hoja
    donde ningun proceso esta en ese estado), se queda con la de mejor
    nombre.
    """
    mejor = None
    for prioridad, indice in candidatas:
        coincidencias = sum(
            1 for fila in filas_datos
            if indice < len(fila) and es_acuerdo_confirmado(fila[indice], exacto)
        )
        clave = (-coincidencias, prioridad)
        if mejor is None or clave < mejor[0]:
            mejor = (clave, indice)
    return mejor[1] if mejor else None


def analizar_hoja(hoja, exacto: bool, columna_forzada=None):
    """
    Devuelve un dict con los datos de la hoja, o None si no tiene una
    columna de estado reconocible.
    """
    filas_crudas = [list(f) for f in hoja.iter_rows(values_only=True)]

    for indice_fila, fila in enumerate(filas_crudas[:MAX_FILAS_ENCABEZADO]):
        if fila_vacia(fila):
            continue
        normalizados = [normalizar(v) for v in fila]

        if columna_forzada:
            objetivo = normalizar(columna_forzada)
            candidatas = [(0, i) for i, e in enumerate(normalizados) if e == objetivo]
        else:
            candidatas = columnas_candidatas(normalizados)
        if not candidatas:
            continue

        datos = [f for f in filas_crudas[indice_fila + 1:] if not fila_vacia(f)]
        columna_estado = elegir_columna_estado(candidatas, datos, exacto)
        if columna_estado is None:
            continue

        encabezados = [limpiar_celda(v) for v in fila]
        while encabezados and (encabezados[-1] is None or str(encabezados[-1]).strip() == ""):
            encabezados.pop()

        filas = []
        for desplazamiento, fila_datos in enumerate(filas_crudas[indice_fila + 1:], start=indice_fila + 2):
            if fila_vacia(fila_datos):
                continue
            estado = fila_datos[columna_estado] if columna_estado < len(fila_datos) else None
            filas.append({
                "fila_original": desplazamiento,
                "estado": estado,
                "excluida": es_acuerdo_confirmado(estado, exacto),
                "valores": [limpiar_celda(v) for v in fila_datos],
            })

        return {
            "titulo": hoja.title,
            "encabezados": encabezados,
            "columna_estado": encabezados[columna_estado] if columna_estado < len(encabezados) else "",
            "filas": filas,
        }

    return None


def ajustar_anchos(hoja, maximo=55):
    anchos = {}
    for fila in hoja.iter_rows(values_only=True):
        for indice, valor in enumerate(fila, start=1):
            if valor is None:
                continue
            largo = max(len(parte) for parte in str(valor).split("\n"))
            anchos[indice] = max(anchos.get(indice, 0), largo)
    for indice, largo in anchos.items():
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(largo + 2, 10), maximo)


AZUL = PatternFill("solid", fgColor="1F3864")
BLANCO_NEGRILLA = Font(color="FFFFFF", bold=True)


def nombre_hoja_valido(titulo, usados):
    limpio = re.sub(r"[\\/*?:\[\]]", "-", str(titulo))[:31] or "Hoja"
    base, contador = limpio, 2
    while limpio.lower() in usados:
        sufijo = f"_{contador}"
        limpio = base[: 31 - len(sufijo)] + sufijo
        contador += 1
    usados.add(limpio.lower())
    return limpio


def escribir_resultado(ruta_salida, hojas, exacto):
    libro = Workbook()
    libro.remove(libro.active)
    resumen = libro.create_sheet("Resumen")
    usados = {"resumen"}

    for datos in hojas:
        hoja = libro.create_sheet(nombre_hoja_valido(datos["titulo"], usados))
        ancho = max([len(datos["encabezados"])] + [len(f["valores"]) for f in datos["filas"]] or [0])
        encabezados = list(datos["encabezados"]) + [None] * (ancho - len(datos["encabezados"]))
        encabezados = [
            e if e not in (None, "") else f"Columna {get_column_letter(i + 1)}"
            for i, e in enumerate(encabezados)
        ]
        hoja.append(["Fila en el informe"] + encabezados)
        for celda in hoja[1]:
            celda.fill = AZUL
            celda.font = BLANCO_NEGRILLA
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.row_dimensions[1].height = 45

        for fila in datos["filas"]:
            if fila["excluida"]:
                continue
            valores = fila["valores"][:ancho]
            valores += [None] * (ancho - len(valores))
            hoja.append([fila["fila_original"]] + valores)

        hoja.freeze_panes = "B2"
        hoja.auto_filter.ref = f"A1:{get_column_letter(ancho + 1)}{hoja.max_row}"
        ajustar_anchos(hoja)

    criterio = "el estado es exactamente ese texto" if exacto else "el estado contiene ese texto"
    resumen.append([f"Procesos que NO estan en estado {ESTADO_A_EXCLUIR}"])
    resumen["A1"].font = Font(bold=True, size=13)
    resumen.append(["Criterio de exclusion", criterio])
    resumen.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    resumen.append([])
    resumen.append(["Hoja", "Columna evaluada", "Procesos", "Excluidos", "En el reporte"])
    for celda in resumen[resumen.max_row]:
        celda.fill = AZUL
        celda.font = BLANCO_NEGRILLA

    total, total_excluidos = 0, 0
    for datos in hojas:
        excluidos = sum(1 for f in datos["filas"] if f["excluida"])
        cantidad = len(datos["filas"])
        total += cantidad
        total_excluidos += excluidos
        resumen.append([datos["titulo"], datos["columna_estado"], cantidad, excluidos, cantidad - excluidos])
    resumen.append(["TOTAL", "", total, total_excluidos, total - total_excluidos])
    for celda in resumen[resumen.max_row]:
        celda.font = Font(bold=True)

    resumen.append([])
    resumen.append(["Detalle por estado"])
    resumen[resumen.max_row][0].font = Font(bold=True)
    resumen.append(["Hoja", "Estado", "Cantidad", "Incluido en el reporte"])
    for celda in resumen[resumen.max_row]:
        celda.fill = AZUL
        celda.font = BLANCO_NEGRILLA

    for datos in hojas:
        conteo = {}
        for fila in datos["filas"]:
            etiqueta = str(fila["estado"]).strip() if fila["estado"] not in (None, "") else "(sin estado)"
            registro = conteo.setdefault(normalizar(etiqueta) or "(SIN ESTADO)",
                                         {"etiqueta": etiqueta, "total": 0, "excluida": fila["excluida"]})
            registro["total"] += 1
        for clave in sorted(conteo, key=lambda k: -conteo[k]["total"]):
            registro = conteo[clave]
            resumen.append([datos["titulo"], registro["etiqueta"], registro["total"],
                            "NO" if registro["excluida"] else "SI"])

    ajustar_anchos(resumen)
    libro.save(ruta_salida)
    return total, total - total_excluidos


def buscar_informe_por_defecto() -> Path:
    carpeta = Path(os.path.expanduser("~")) / "Downloads"
    candidatos = sorted(
        (p for p in carpeta.glob("*.xlsx")
         if p.name.lower().startswith("informe procesos concursales")
         and not p.name.startswith("~$")
         and "SIN ACUERDO CONFIRMADO" not in p.name.upper()),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidatos:
        raise SystemExit(
            f"No se encontro ningun 'informe procesos concursales*.xlsx' en {carpeta}.\n"
            f"Pasa la ruta del archivo como argumento:\n"
            f'    python filtrar_procesos.py "C:\\ruta\\al\\informe.xlsx"'
        )
    return candidatos[0]


def main():
    parser = argparse.ArgumentParser(
        description="Extrae los procesos que NO estan en estado 'ACUERDO CONFIRMADO' y los deja en un Excel nuevo."
    )
    parser.add_argument("informe", nargs="?", help="Ruta del .xlsx del informe de procesos concursales.")
    parser.add_argument("-o", "--salida", help="Ruta del Excel de salida.")
    parser.add_argument("--columna", help="Nombre exacto de la columna de estado, si se quiere forzar "
                                          "(por ejemplo: \"ETAPA ACTUAL\").")
    parser.add_argument(
        "--exacto",
        action="store_true",
        help="Excluir solo los estados que digan exactamente 'ACUERDO CONFIRMADO' "
             "(por defecto tambien se excluyen variantes como 'Acuerdo confirmado en ejecucion').",
    )
    args = parser.parse_args()

    ruta_informe = Path(args.informe).expanduser() if args.informe else buscar_informe_por_defecto()
    if not ruta_informe.is_file():
        raise SystemExit(f"No existe el archivo: {ruta_informe}")
    if ruta_informe.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit(
            f"'{ruta_informe.name}' no es un archivo .xlsx.\n"
            f"Abrelo en Excel y guardalo como 'Libro de Excel (*.xlsx)' antes de correr este programa."
        )

    print(f"Leyendo: {ruta_informe}")
    libro = load_workbook(ruta_informe, data_only=True)
    hojas = []
    for hoja in libro.worksheets:
        datos = analizar_hoja(hoja, args.exacto, args.columna)
        if datos is None:
            print(f"  (aviso) la hoja '{hoja.title}' no tiene columna de estado, se omitio.")
            continue
        hojas.append(datos)
        print(f"  hoja '{hoja.title}': columna de estado = '{datos['columna_estado']}'")
    libro.close()

    if not hojas:
        columnas = ", ".join(NOMBRES_COLUMNA_ESTADO)
        raise SystemExit(
            f"No se encontro ninguna columna de estado en '{ruta_informe.name}'.\n"
            f"Se buscaron columnas llamadas: {columnas}.\n"
            f'Puedes indicarla a mano con: --columna "NOMBRE DE LA COLUMNA"'
        )

    if args.salida:
        ruta_salida = Path(args.salida).expanduser()
    else:
        ruta_salida = ruta_informe.with_name(f"{ruta_informe.stem} - SIN ACUERDO CONFIRMADO.xlsx")

    total, incluidos = escribir_resultado(ruta_salida, hojas, args.exacto)

    print()
    for datos in hojas:
        excluidos = sum(1 for f in datos["filas"] if f["excluida"])
        print(f"  {datos['titulo']}: {len(datos['filas']) - excluidos} de {len(datos['filas'])} procesos")
    print()
    print(f"Procesos en el informe : {total}")
    print(f"Sin acuerdo confirmado : {incluidos}")
    print(f"Excel generado         : {ruta_salida}")


if __name__ == "__main__":
    sys.exit(main())
