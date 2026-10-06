"""Pruebas de construir_bd.py con archivos de ejemplo (no usa internet).

Ejecutar con:  python -m unittest test_construir_bd
"""

import contextlib
import csv
import io
import sqlite3
import tempfile
import unittest
from unittest import mock
from pathlib import Path

import construir_bd as bd

REPS = [
    ["CodigoPrestador", "NombrePrestador", "CodigoHabilitacionSede",
     "NombreSede", "SedePrincipal", "TipoIdentificacion",
     "NumeroIdentificacion", "ClasePrestador", "NaturalezaJuridica",
     "DepartamentoPrestadorDesc", "MunicipioPrestadorDesc",
     "DepartamentoSedeDesc", "MunicipioSedeDesc", "Direccion", "Telefono"],
    # IPS con dos sedes en Antioquia: debe quedar una sola fila con 2 sedes.
    ["0500100001", "CLINICA SAN JUAN SAS", "050010000102", "SEDE ENVIGADO",
     "NO", "NI", "900111222", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Antioquia", "MEDELLÍN", "Antioquia", "ENVIGADO", "CL 1", "111"],
    ["0500100001", "CLINICA SAN JUAN SAS", "050010000101", "SEDE PRINCIPAL",
     "SI", "NI", "900111222", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Antioquia", "MEDELLÍN", "Antioquia", "MEDELLÍN", "CL 2", "222"],
    # Norte de Santander escrito abreviado.
    ["5400100002", "IPS FRONTERA LTDA", "540010000201", "IPS FRONTERA LTDA",
     "SI", "NI", "800222333-4", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "N. de Santander", "CÚCUTA", "N. de Santander", "CÚCUTA", "AV 0", "333"],
    # Santander no es Norte de Santander.
    ["6800100003", "IPS BUCARAMANGA", "680010000301", "IPS BUCARAMANGA",
     "SI", "NI", "900333444", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Santander", "BUCARAMANGA", "Santander", "BUCARAMANGA", "CR 1", "444"],
    # Bogota solo entra con --incluir-bogota.
    ["1100100004", "IPS CAPITAL", "110010000401", "IPS CAPITAL",
     "SI", "NI", "900444555", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Bogotá D.C.", "BOGOTÁ", "Bogotá D.C.", "BOGOTÁ", "CR 7", "555"],
    # Prestador con domicilio en Cundinamarca pero sede en Bogota: la sede
    # esta en Bogota.
    ["2575400005", "IPS SOACHA", "257540000502", "SEDE BOGOTA",
     "NO", "NI", "900555666", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Cundinamarca", "SOACHA", "Bogotá D.C.", "BOGOTÁ", "CL 80", "666"],
    # Profesional independiente: va a reps_todas_las_sedes pero no es IPS.
    ["0800100006", "PEREZ GOMEZ JUAN", "080010000601", "CONSULTORIO",
     "SI", "CC", "72123456", "Profesional Independiente",
     "Privada", "Atlántico", "BARRANQUILLA", "Atlántico", "BARRANQUILLA", "CL 3", "777"],
    # IPS cuyo NIT (con digito de verificacion) es de una cooperativa de la
    # Supersolidaria registrada en otro departamento.
    ["1300100007", "SERVISALUD BOLIVAR", "130010000701", "SERVISALUD",
     "SI", "NI", "890.123.456-7", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Bolívar", "CARTAGENA", "Bolívar", "CARTAGENA", "CL 4", "888"],
    # La misma IPS (mismo NIT) inscrita en dos departamentos con codigos de
    # prestador distintos: debe quedar una sola vez en la hoja IPS.
    ["0800103693", "DAVITA S.A.S.", "080010369301", "DAVITA BARRANQUILLA",
     "SI", "NI", "900532504", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Atlántico", "BARRANQUILLA", "Atlántico", "BARRANQUILLA", "CL 6", "101"],
    ["0800103693", "DAVITA S.A.S.", "080010369302", "DAVITA NORTE",
     "NO", "NI", "900532504", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Atlántico", "BARRANQUILLA", "Atlántico", "BARRANQUILLA", "CL 7", "102"],
    ["1300103270", "DAVITA S.A.S.", "130010327001", "DAVITA CARTAGENA",
     "SI", "NI", "900532504", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Bolívar", "CARTAGENA", "Bolívar", "CARTAGENA", "CL 8", "103"],
    # IPS que es cooperativa por su razon social.
    ["2530700008", "COOPERATIVA DE SALUD DE GIRARDOT", "253070000801", "COOPSALUD",
     "SI", "NI", "900777888", "Instituciones Prestadoras de Servicios de Salud - IPS",
     "Privada", "Cundinamarca", "GIRARDOT", "Cundinamarca", "GIRARDOT", "CL 5", "999"],
]

SOLIDARIAS = [
    # Mismo formato que el listado real: un registro por cada reporte.
    ["codentidad", "fechaultirepo", "mes", "ano", "nombreentidad", "sigla", "nit",
     "nombretipo", "departamento", "municipio", "direccion", "supervision"],
    # Dos reportes de la misma cooperativa: queda el mas reciente.
    ["1", "12/31/2025 05:00:00 PM", "DICIEMBRE", "2025", "COOPERATIVA VIEJA", "",
     "890-900-111-1", "Especializada de ahorro y credito", "ANTIOQUIA", "MEDELLÍN",
     "CL 10", "1"],
    ["1", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "COOPERATIVA ANTIOQUEÑA", "COOPANT",
     "890-900-111-1", "Especializada de ahorro y credito", "ANTIOQUIA", "MEDELLÍN",
     "CL 10", "1"],
    # Cooperativa cuyo nombre tiene un error de digitacion: cuenta por su tipo.
    ["2", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "COOPERATVA DE MILITARES", "",
     "800-100-600-6", "Multiactiva sin seccion de ahorro", "ANTIOQUIA", "BELLO",
     "CL 16", "3"],
    # Cooperativa que dejo de reportar hace anios.
    ["3", "12/31/2019 05:00:00 PM", "DICIEMBRE", "2019", "COOPERATIVA INACTIVA", "",
     "800-100-700-7", "Multiactiva sin seccion de ahorro", "CUNDINAMARCA", "CHÍA",
     "CL 17", "3"],
    # Departamento con la codificacion danada, como llega en el listado real.
    ["4", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "FONDO DE EMPLEADOS DEL CARIBE",
     "FONCARIBE", "800-100-200-2", "Fondos de empleados", "ATLÃ\x81NTICO",
     "BARRANQUILLA", "CL 11", "2"],
    ["5", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "INSTITUTO DEL COOPERATIVISMO", "",
     "800-100-300-3", "Instituciones auxiliares especializadas", "BOLÃ\x8dVAR",
     "CARTAGENA", "CL 12", "3"],
    ["6", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "PRECOOPERATIVA DEL MAR", "",
     "800-100-400-4", "Precooperativas", "BOLÃ\x8dVAR", "CARTAGENA", "CL 13", "3"],
    ["7", "07/31/2026 05:00:00 PM", "JULIO", "2026",
     "COOPERATIVA DE TRABAJO ASOCIADO NORTE", "", "800-100-500-5",
     "Cooperativas de trabajo asociado", "NORTE DE SANTANDER", "CÚCUTA", "CL 14", "1"],
    # Cooperativa fuera de los departamentos; su NIT marca a la IPS de Bolivar.
    ["8", "06/30/2026 05:00:00 PM", "JUNIO", "2026", "COOPERATIVA DEL VALLE", "",
     "890-123-456-7", "Multiactiva con ahorro y credito", "VALLE DEL CAUCA", "CALI",
     "CL 15", "1"],
]


# Mismo formato que la capacidad instalada real (s2ru-bqt6). El codigo de la
# sede empieza por el codigo DANE de su municipio (sin el cero inicial).
CAPACIDAD = [
    ["Departamento", "Municipio", "Código prestador", "Nombre prestador", "nit IPS ",
     "num digito_verificion", "naturaleza", "num nivel atencion", "Código sede",
     "Número sede", "nom sede IPS", "nom grupo capacidad ",
     "nom descripcion capacidad ", "num cantidad capacidad instalada", "Fecha Corte"],
    ["Antioquia", "MEDELLÍN", "500100001", "CLINICA SAN JUAN SAS", "900111222", "1",
     "Privada", "", "500100001", "01", "SEDE PRINCIPAL", "CAMAS", "Adultos", "100",
     "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Antioquia", "MEDELLÍN", "500100001", "CLINICA SAN JUAN SAS", "900111222", "1",
     "Privada", "", "500100001", "01", "SEDE PRINCIPAL", "SALAS", "Sala de Cirugía", "5",
     "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Antioquia", "ENVIGADO", "500100001", "CLINICA SAN JUAN SAS", "900111222", "1",
     "Privada", "", "526600001", "02", "SEDE ENVIGADO", "CONSULTORIOS",
     "Consulta Externa", "20", "Fecha corte REPS: Nov  5 2022  1:37PM"],
    # Davita: sillas de hemodialisis en Barranquilla y Cartagena; las de Cali
    # no cuentan porque no estan en los departamentos.
    ["Barranquilla", "BARRANQUILLA", "800103693", "DAVITA S.A.S.", "900532504", "4",
     "Privada", "", "800103693", "01", "DAVITA BARRANQUILLA", "SILLAS",
     "Sillas de Hemodiálisis", "40", "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Cartagena", "CARTAGENA", "1300103270", "DAVITA S.A.S.", "900532504", "4",
     "Privada", "", "1300103270", "01", "DAVITA CARTAGENA", "SILLAS",
     "Sillas de Hemodiálisis", "20", "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Cali", "CALI", "7600100001", "DAVITA S.A.S.", "900532504", "4",
     "Privada", "", "7600100001", "01", "DAVITA CALI", "SILLAS",
     "Sillas de Hemodiálisis", "500", "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Bolívar", "CARTAGENA", "1300100007", "SERVISALUD BOLIVAR", "890123456", "7",
     "Pública", "2", "1300100007", "01", "SERVISALUD", "CONSULTORIOS",
     "Consulta Externa", "3", "Fecha corte REPS: Nov  5 2022  1:37PM"],
    ["Cundinamarca", "GIRARDOT", "2530700008", "COOPERATIVA DE SALUD DE GIRARDOT",
     "900777888", "1", "Privada", "", "2530700008", "01", "COOPSALUD", "CAMAS",
     "Pediátrica", "10", "Fecha corte REPS: Nov  5 2022  1:37PM"],
]

# Mismo formato que los estados financieros reales (tic6-rbue).
ACTIVOS = [
    ["AÑO", "MES", "CODIGO ENTIDAD", "NIT", "CODRENGLON", "NOMBRE CUENTA", "VALOR EN PESOS"],
    ["2025", "DICIEMBRE", "1", "890-900-111-1", "100000", "ACTIVO", "$    4,000,000,000.00"],
    ["2026", "JUNIO", "1", "890-900-111-1", "100000", "ACTIVO", "$    5,000,000,000.00"],
    ["2026", "JULIO", "7", "800-100-500-5", "100000", "ACTIVO", "$    80,000,000,000.00"],
    # Otra cuenta: no es el total de activos.
    ["2026", "JULIO", "7", "800-100-500-5", "110000", "EFECTIVO", "$    999,000,000,000,000.00"],
    ["2026", "JUNIO", "2", "800-100-600-6", "100000", "ACTIVO", "1.000.000"],
]


def escribir_csv(ruta, filas):
    with open(ruta, "w", encoding="utf-8", newline="") as archivo:
        csv.writer(archivo).writerows(filas)


class PruebasTexto(unittest.TestCase):
    def test_departamento_de(self):
        objetivos = bd.construir_objetivos()
        casos = {
            "Antioquia": "ANTIOQUIA",
            "ATLANTICO": "ATLÁNTICO",
            "Bolívar": "BOLÍVAR",
            "N. de Santander": "NORTE DE SANTANDER",
            "NORTE SANTANDER": "NORTE DE SANTANDER",
            "Departamento de Cundinamarca": "CUNDINAMARCA",
            "05 - ANTIOQUIA": "ANTIOQUIA",
            "54": "NORTE DE SANTANDER",
            "5": "ANTIOQUIA",
            "13001": "BOLÍVAR",
            "5001": "ANTIOQUIA",
            "Santander": None,
            "Bogotá D.C.": None,
            "Valle del Cauca": None,
            "": None,
        }
        for valor, esperado in casos.items():
            self.assertEqual(bd.departamento_de(valor, objetivos), esperado, valor)
        con_bogota = bd.construir_objetivos(incluir_bogota=True)
        self.assertEqual(bd.departamento_de("BOGOTA DC", con_bogota), "BOGOTÁ D.C.")
        self.assertEqual(bd.departamento_de("11", con_bogota), "BOGOTÁ D.C.")

    def test_a_numero(self):
        casos = {"$    81,781,823,706.19": 81781823706.19, "1.000.000": 1000000,
                 "0,00": 0, "$    - 0": 0, "12,5": 12.5, "": 0}
        for valor, esperado in casos.items():
            self.assertEqual(bd.a_numero(valor), esperado, valor)

    def test_normalizar_nit(self):
        for valor in ("890.123.456-7", "890-123-456-7", "8901234567", "890123456",
                      "890123456.0", "890-123-456"):
            self.assertEqual(bd.normalizar_nit(valor), "890123456", valor)
        self.assertEqual(bd.normalizar_nit("72123456"), "72123456")
        self.assertEqual(bd.normalizar_nit(""), "")

    def test_patron_cooperativa(self):
        for texto in ("COOPERATIVA DE AHORRO", "PRECOOPERATIVA", "COOP. MEDICA",
                      "ORGANISMO COOPERATIVO DE GRADO SUPERIOR"):
            self.assertTrue(bd.PATRON_COOPERATIVA.search(texto), texto)
        for texto in ("INSTITUCION AUXILIAR DEL COOPERATIVISMO", "COOPSALUD",
                      "FONDO DE EMPLEADOS"):
            self.assertFalse(bd.PATRON_COOPERATIVA.search(texto), texto)

    def test_reparar_texto(self):
        self.assertEqual(bd.reparar_texto("BOLÃ\x8dVAR"), "BOLÍVAR")
        self.assertEqual(bd.reparar_texto("ATLÃ\x81NTICO"), "ATLÁNTICO")
        self.assertEqual(bd.reparar_texto("NARIÃ‘O"), "NARIÑO")
        self.assertEqual(bd.reparar_texto("BOGOTÁ, D.C."), "BOGOTÁ, D.C.")

    def test_extraer_correos(self):
        casos = {
            "calidad@davita.com   legaldavita@davita.com":
                [("calidad@davita.com", ""), ("legaldavita@davita.com", "")],
            "a@viva1a.com.co - b@viva1a.com.co -":
                [("a@viva1a.com.co", ""), ("b@viva1a.com.co", "")],
            "ceginob@hotmail.com-ceginob@gmail.com":
                [("ceginob@hotmail.com", ""), ("ceginob@gmail.com", "")],
            "gerencia.usos@ gmail.com": [("gerencia.usos@gmail.com", "")],
            "direccióncalidad@colcan.com":
                [("direccioncalidad@colcan.com", "se le quitó una tilde")],
            "Y@GMAIL.CON": [("y@gmail.com", "dominio corregido, decía gmail.con")],
            "x@gmail.c": [("x@gmail.com", "dominio corregido, decía gmail.c")],
            "a@x.com. b@y.com.co.": [("a@x.com", ""), ("b@y.com.co", "")],
            "a@x.com, a@x.com": [("a@x.com", "")],
            "gerencia@cediul": [],
            "www.hospital.gov.co": [],
            "gerencia@hospital de calamar.gov.co": [],
        }
        for valor, esperado in casos.items():
            self.assertEqual(bd.extraer_correos(valor), esperado, valor)

    def test_nombre_columna(self):
        self.assertEqual(bd.nombre_columna("NombrePrestador"), "nombre_prestador")
        self.assertEqual(bd.nombre_columna("RAZÓN SOCIAL"), "razon_social")
        self.assertEqual(bd.nombre_columna("DepartamentoSedeDesc"),
                         "departamento_sede_desc")
        self.assertEqual(bd.unicos(["a", "a", "a_2", "b"]), ["a", "a_2", "a_2_2", "b"])

    def test_clave_fecha(self):
        self.assertGreater(bd.clave_fecha("08/31/2026 12:00:00 AM"),
                           bd.clave_fecha("12/31/2025 12:00:00 AM"))
        self.assertEqual(bd.clave_fecha("2026-08-31T00:00:00.000"), (2026, 8, 31))
        self.assertEqual(bd.clave_fecha("31/08/2026"), (2026, 8, 31))
        self.assertEqual(bd.clave_fecha("202608"), (2026, 8, 0))
        self.assertEqual(bd.clave_fecha("Fecha corte REPS: Mar 12 2026  3:11PM"),
                         (2026, 3, 12))
        self.assertEqual(bd.clave_fecha("Fecha corte REPS: Nov  5 2022  1:37PM"),
                         (2022, 11, 5))
        self.assertEqual(bd.restar_meses((2026, 7, 31), 12), (2025, 7, 31))


class PruebaCompleta(unittest.TestCase):
    def construir(self, *extra):
        carpeta = Path(self.enterContext(tempfile.TemporaryDirectory()))
        escribir_csv(carpeta / "reps.csv", REPS)
        escribir_csv(carpeta / "solidarias.csv", SOLIDARIAS)
        escribir_csv(carpeta / "capacidad.csv", CAPACIDAD)
        escribir_csv(carpeta / "activos.csv", ACTIVOS)
        salida = carpeta / "salida"
        tamano = ["--archivo-capacidad", str(carpeta / "capacidad.csv"),
                  "--archivo-activos", str(carpeta / "activos.csv")]
        with contextlib.redirect_stdout(io.StringIO()):
            codigo = bd.main(["--archivo-reps", str(carpeta / "reps.csv"),
                              "--archivo-solidarias", str(carpeta / "solidarias.csv"),
                              "--salida", str(salida), *(extra or tamano)])
        self.assertEqual(codigo, 0)
        con = sqlite3.connect(salida / bd.NOMBRE_SQLITE)
        self.addCleanup(con.close)
        return salida, con

    def test_ips(self):
        _, con = self.construir()
        filas = con.execute(
            "SELECT bd_departamento, bd_municipio, nombre_prestador, bd_num_sedes, "
            "nombre_sede, bd_es_cooperativa, bd_criterio_cooperativa FROM ips "
            "ORDER BY bd_departamento, nombre_prestador").fetchall()
        self.assertEqual(filas, [
            ("ANTIOQUIA", "MEDELLIN", "CLINICA SAN JUAN SAS", 2, "SEDE PRINCIPAL", "NO", ""),
            ("ATLÁNTICO", "BARRANQUILLA", "DAVITA S.A.S.", 3, "DAVITA BARRANQUILLA",
             "NO", ""),
            ("BOLÍVAR", "CARTAGENA", "SERVISALUD BOLIVAR", 1, "SERVISALUD", "SI",
             "NIT registrado como cooperativa en la Supersolidaria"),
            ("CUNDINAMARCA", "GIRARDOT", "COOPERATIVA DE SALUD DE GIRARDOT", 1,
             "COOPSALUD", "SI", "Razón social de cooperativa"),
            ("NORTE DE SANTANDER", "CUCUTA", "IPS FRONTERA LTDA", 1, "IPS FRONTERA LTDA",
             "NO", ""),
        ])
        self.assertEqual(con.execute(
            "SELECT bd_departamentos, bd_inscripciones_reps FROM ips "
            "WHERE nombre_prestador = 'DAVITA S.A.S.'").fetchall(),
            [("ATLÁNTICO, BOLÍVAR", 2)])
        self.assertEqual(con.execute("SELECT COUNT(*) FROM ips_sedes").fetchone(), (8,))
        self.assertEqual(con.execute("SELECT COUNT(*) FROM ips_cooperativas").fetchone(), (2,))
        # El profesional independiente queda en el REPS completo, no como IPS.
        self.assertEqual(con.execute(
            "SELECT bd_departamento, bd_es_ips FROM reps_todas_las_sedes "
            "WHERE nombre_prestador = 'PEREZ GOMEZ JUAN'").fetchall(),
            [("ATLÁNTICO", "NO")])

    def test_cooperativas(self):
        _, con = self.construir()
        filas = con.execute(
            "SELECT bd_departamento, bd_municipio, nombreentidad, bd_reporta_actualmente, "
            "bd_ultimo_reporte FROM cooperativas ORDER BY bd_departamento, "
            "nombreentidad").fetchall()
        self.assertEqual(filas, [
            ("ANTIOQUIA", "MEDELLIN", "COOPERATIVA ANTIOQUEÑA", "SI", "2026-06-30"),
            ("ANTIOQUIA", "BELLO", "COOPERATVA DE MILITARES", "SI", "2026-06-30"),
            ("BOLÍVAR", "CARTAGENA", "PRECOOPERATIVA DEL MAR", "SI", "2026-06-30"),
            ("CUNDINAMARCA", "CHIA", "COOPERATIVA INACTIVA", "NO", "2019-12-31"),
            ("NORTE DE SANTANDER", "CUCUTA", "COOPERATIVA DE TRABAJO ASOCIADO NORTE",
             "SI", "2026-07-31"),
        ])
        self.assertEqual(con.execute(
            "SELECT bd_departamento, nombreentidad FROM entidades_solidarias "
            "WHERE bd_es_cooperativa = 'NO' ORDER BY nombreentidad").fetchall(), [
            ("ATLÁNTICO", "FONDO DE EMPLEADOS DEL CARIBE"),
            ("BOLÍVAR", "INSTITUTO DEL COOPERATIVISMO"),
        ])
        self.assertEqual(con.execute(
            "SELECT dataset, corte_de_los_datos FROM fuentes ORDER BY dataset").fetchall(),
            [("c36g-9fc2", ""), ("kg2d-yfyg", "2026-07-31"), ("s2ru-bqt6", "2022-11-05"),
             ("tic6-rbue", "2026-07")])

    def test_resumen_y_archivos(self):
        salida, con = self.construir()
        total = con.execute(
            "SELECT ips, sedes_de_ips, ips_cooperativas, otras_sedes_reps_no_ips, "
            "cooperativas, cooperativas_que_reportan_actualmente, "
            "otras_entidades_solidarias FROM resumen "
            "WHERE departamento = 'TOTAL'").fetchone()
        # Davita cuenta en Atlantico y en Bolivar, pero una sola vez en el total.
        self.assertEqual(total, (5, 8, 2, 1, 5, 4, 2))
        self.assertEqual(con.execute(
            "SELECT departamento, ips FROM resumen WHERE departamento IN "
            "('ATLÁNTICO', 'BOLÍVAR') ORDER BY 1").fetchall(),
            [("ATLÁNTICO", 1), ("BOLÍVAR", 2)])

        from openpyxl import load_workbook
        libro = load_workbook(salida / bd.NOMBRE_EXCEL)
        self.assertEqual(libro.sheetnames, ["Resumen", "IPS", "Sedes IPS",
                                            "Cooperativas", "IPS cooperativas",
                                            "Fuentes"])
        hoja = libro["IPS"]
        self.assertEqual(hoja.max_row, 6)  # 5 IPS + encabezado
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertTrue(hoja.auto_filter.ref)
        self.assertEqual(hoja["A1"].value, "bd_ranking_tamano")
        self.assertTrue((salida / "csv" / "ips.csv").exists())
        correos = load_workbook(salida / bd.NOMBRE_CORREOS)
        self.assertEqual(correos.sheetnames,
                         ["Correos", "Correos IPS cooperativas", "Por revisar"])

    def test_sede_principal_por_codigo(self):
        # El REPS real no marca la sede principal: es la del codigo + "01".
        encabezados = ["codigo_prestador", "codigo_habilitacion_sede", "nombre_sede"]
        sedes = [
            {"codigo_prestador": "0523704806", "codigo_habilitacion_sede": cod,
             "nombre_sede": nombre, "bd_es_ips": "SI", "bd_departamento": "ANTIOQUIA",
             "bd_es_cooperativa": "NO", "bd_criterio_cooperativa": ""}
            for cod, nombre in (("052370480607", "SEDE BARBOSA"),
                                ("052370480601", "SEDE DONMATIAS"),
                                ("052370480603", "SEDE GIRARDOTA"))]
        ips = bd.agrupar_ips(encabezados, sedes)
        self.assertEqual([(i["nombre_sede"], i["bd_num_sedes"]) for i in ips],
                         [("SEDE DONMATIAS", 3)])

    def test_lista_de_correos(self):
        encabezados = ["numero_identificacion", "nombre_prestador",
                       "email_prestador", "email_sede"]

        def sede(nit, nombre, depto, muni, principal, de_sede, coop="NO"):
            return {"numero_identificacion": nit, "nombre_prestador": nombre,
                    "email_prestador": principal, "email_sede": de_sede,
                    "bd_departamento": depto, "bd_municipio": muni,
                    "bd_es_cooperativa": coop}

        sedes = [
            # Misma IPS en dos departamentos: su correo principal sale una vez.
            sede("900532504", "DAVITA S.A.S.", "ATLÁNTICO", "BARRANQUILLA",
                 "calidad@davita.com", "baq@davita.com"),
            sede("900532504", "DAVITA S.A.S.", "BOLÍVAR", "CARTAGENA",
                 "calidad@davita.com", "ctg@davita.com"),
            sede("802007499", "COOPERATIVA CONSALUD", "ATLÁNTICO", "SOLEDAD",
                 "consalud@gmail.con", "consalud@gmail.com", coop="SI"),
            sede("900111222", "IPS SIN CORREO", "BOLÍVAR", "MAGANGUE",
                 "www.ips.com", "gerencia@ips"),
        ]
        correos, revisar = bd.lista_de_correos(encabezados, sedes)
        por_correo = {c["Correo"]: c for c in correos}
        self.assertEqual(sorted(por_correo), [
            "baq@davita.com", "calidad@davita.com", "consalud@gmail.com",
            "ctg@davita.com"])
        davita = por_correo["calidad@davita.com"]
        self.assertEqual((davita["IPS"], davita["Departamentos"], davita["Tipo de correo"]),
                         ("DAVITA S.A.S.", "ATLÁNTICO, BOLÍVAR", "Principal de la IPS"))
        self.assertEqual(por_correo["baq@davita.com"]["Tipo de correo"], "De una sede")
        consalud = por_correo["consalud@gmail.com"]
        self.assertEqual((consalud["Es cooperativa"], consalud["Observación"]),
                         ("SI", "dominio corregido, decía gmail.con"))
        self.assertEqual([(r["IPS"], r["Valor en el REPS"], r["La IPS tiene otro correo válido"])
                          for r in revisar],
                         [("IPS SIN CORREO", "www.ips.com", "NO"),
                          ("IPS SIN CORREO", "gerencia@ips", "NO")])

    def test_orden_por_tamano(self):
        salida, con = self.construir()
        self.assertEqual(con.execute(
            "SELECT bd_ranking_tamano, nombre_prestador, bd_capacidad_instalada, bd_camas, "
            "bd_consultorios, bd_salas, bd_nivel_atencion FROM ips").fetchall(), [
            (1, "CLINICA SAN JUAN SAS", 125, 100, 20, 5, ""),
            (2, "DAVITA S.A.S.", 60, 0, 0, 0, ""),
            (3, "COOPERATIVA DE SALUD DE GIRARDOT", 10, 10, 0, 0, ""),
            (4, "SERVISALUD BOLIVAR", 3, 0, 3, 0, "2"),
            (5, "IPS FRONTERA LTDA", "", "", "", "", ""),
        ])
        # Las sedes siguen el orden de su IPS.
        self.assertEqual([n for (n,) in con.execute(
            "SELECT DISTINCT nombre_prestador FROM ips_sedes")][:2],
            ["CLINICA SAN JUAN SAS", "DAVITA S.A.S."])
        self.assertEqual(con.execute(
            "SELECT bd_ranking_tamano, nombreentidad, bd_activos_pesos, bd_fecha_activos "
            "FROM cooperativas").fetchall(), [
            (1, "COOPERATIVA DE TRABAJO ASOCIADO NORTE", 80000000000, "2026-07"),
            (2, "COOPERATIVA ANTIOQUEÑA", 5000000000, "2026-06"),
            (3, "COOPERATVA DE MILITARES", 1000000, "2026-06"),
            (4, "COOPERATIVA INACTIVA", "", ""),
            (5, "PRECOOPERATIVA DEL MAR", "", ""),
        ])
        from openpyxl import load_workbook
        correos = load_workbook(salida / bd.NOMBRE_CORREOS)["Correos"]
        self.assertEqual(correos["A1"].value, "Ranking tamaño")

    def test_sin_datos_de_tamano(self):
        # Si no hay capacidad ni activos, igual se arma la base: las IPS se
        # ordenan por numero de sedes.
        _, con = self.construir("--archivo-capacidad", "no_existe.csv",
                                "--archivo-activos", "no_existe.csv")
        self.assertEqual(con.execute(
            "SELECT nombre_prestador, bd_num_sedes FROM ips "
            "WHERE bd_ranking_tamano <= 2").fetchall(),
            [("DAVITA S.A.S.", 3), ("CLINICA SAN JUAN SAS", 2)])
        self.assertEqual(con.execute("SELECT COUNT(*) FROM fuentes").fetchone(), (2,))

    def test_obtener_activos_por_nit(self):
        # Se consulta NIT por NIT; si uno falla, se siguen los demas.
        respuestas = {
            "890-900-111-1": [{"a_o": "2026", "mes": "JUNIO", "nit": "890-900-111-1",
                               "valor_en_pesos": "$    5,000,000.00"}],
            "800-100-500-5": [],
        }

        def consultar(dataset, consulta, app_token=None):
            nit = consulta["$where"].split("'")[1]
            if nit not in respuestas:
                raise OSError("503")
            return respuestas[nit]

        carpeta = Path(self.enterContext(tempfile.TemporaryDirectory()))
        with mock.patch.object(bd, "consultar", consultar), \
                contextlib.redirect_stdout(io.StringIO()):
            ruta, origen = bd.obtener_activos(
                None, carpeta, None, ["890-900-111-1", "800-100-500-5", "800-100-600-6"])
        cols, filas = bd.leer_tabla(ruta)
        self.assertEqual(bd.activos_por_nit(cols, filas),
                         {"890900111": ((2026, 6), 5000000.0)})
        self.assertIn("2 entidades", origen)

        # Si fallan todas y no hay copia anterior, no hay activos.
        with mock.patch.object(bd, "consultar", consultar), \
                contextlib.redirect_stdout(io.StringIO()), \
                self.assertRaises(RuntimeError):
            bd.obtener_activos(None, carpeta / "vacia", None, ["800-100-600-6"])

    def test_incluir_bogota(self):
        _, con = self.construir("--incluir-bogota", "--archivo-capacidad", "no_existe.csv",
                                "--archivo-activos", "no_existe.csv")
        self.assertEqual(con.execute(
            "SELECT nombre_prestador FROM ips WHERE bd_departamento = 'BOGOTÁ D.C.' "
            "ORDER BY nombre_prestador").fetchall(),
            [("IPS CAPITAL",), ("IPS SOACHA",)])

    def test_lee_xlsx(self):
        from openpyxl import Workbook
        carpeta = Path(self.enterContext(tempfile.TemporaryDirectory()))
        libro = Workbook()
        libro.active.append(["Listado de prestadores"])  # titulo antes de la tabla
        for fila in REPS:
            libro.active.append(fila)
        libro.save(carpeta / "reps.xlsx")
        cols, filas = bd.leer_tabla(carpeta / "reps.xlsx")
        self.assertEqual(cols[:2], ["codigo_prestador", "nombre_prestador"])
        self.assertEqual(len(filas), len(REPS) - 1)


if __name__ == "__main__":
    unittest.main()
