"""
Script individual y de SOLO LECTURA: compara, por RADICADO, los
procesos del informe de Excel contra las carpetas que existen en el
disco duro, y genera un reporte de cuales procesos del Excel NO tienen
ninguna carpeta correspondiente en el disco.

Dos radicados se consideran el MISMO proceso si son idénticos, o si
tienen los primeros 22 digitos iguales y solo difieren en el ultimo
(el "consecutivo" de instancia/reparto -- ej. termina en 0 o en 1); eso
no es un proceso distinto (misma logica que
mismo_radicado_salvo_ultimo_digito en validar_renombrar_carpetas.py).

Lee el Excel de forma independiente (solo necesita la columna
RADICADO) -- no depende de ninguna columna de numero de proceso, asi
que funciona con informes que no tengan una columna "No.".

Este script NUNCA mueve, renombra, crea ni borra nada -- solo compara
y genera un reporte (comparar_excel_disco_faltantes.csv) con los
procesos del Excel que no tienen carpeta en el disco, mas el log
detallado (comparar_excel_disco.log).

Usa la misma configuracion (RUTA_EXCEL, HOJA_EXCEL, CARPETA_PROCESOS,
etc) de validar_renombrar_carpetas.py -- no hay que configurarla dos
veces.
"""

import logging
import os
import re
from pathlib import Path

import openpyxl

import validar_renombrar_carpetas as cruce_excel

# ============================= CONFIGURACION =============================

ARCHIVO_LOG = os.path.join(os.path.dirname(__file__), "comparar_excel_disco.log")
ARCHIVO_REPORTE = os.path.join(os.path.dirname(__file__), "comparar_excel_disco_faltantes.csv")

# ===========================================================================


def configurar_logging():
    logging.basicConfig(
        level=logging.INFO,
        format="%(message)s",
        handlers=[
            logging.FileHandler(ARCHIVO_LOG, encoding="utf-8"),
            logging.StreamHandler(),
        ],
    )


def _radicados_en_disco():
    """
    Devuelve (exactos, bases): 'exactos' es el set de radicados de 23
    digitos encontrados en el nombre de las carpetas de primer nivel del
    disco; 'bases' es el set de sus primeros 22 digitos (para poder
    reconocer un radicado que solo cambia en el ultimo digito).
    """
    carpeta_raiz = Path(cruce_excel.CARPETA_PROCESOS)
    if not carpeta_raiz.exists():
        logging.error("[Disco] No existe la carpeta configurada en CARPETA_PROCESOS: %s", carpeta_raiz)
        return set(), set()

    try:
        carpetas = [
            d for d in carpeta_raiz.iterdir()
            if d.is_dir() and d.name not in cruce_excel.CARPETAS_A_IGNORAR
        ]
    except OSError as error:
        logging.error("[Disco] No se pudo leer %s: %s", carpeta_raiz, error)
        return set(), set()

    exactos = set()
    for carpeta in carpetas:
        radicado = cruce_excel.radicado_de_nombre_carpeta(carpeta.name)
        if radicado:
            exactos.add(radicado)

    bases = {radicado[:-1] for radicado in exactos}
    return exactos, bases


def leer_radicados_excel():
    """
    Lee el Excel y devuelve [(fila, radicado), ...] para cada fila con un
    radicado de EXACTAMENTE 23 digitos. Solo necesita la columna RADICADO
    -- no depende de ninguna columna de numero de proceso.
    """
    wb = openpyxl.load_workbook(cruce_excel.RUTA_EXCEL, data_only=True)
    if cruce_excel.HOJA_EXCEL not in wb.sheetnames:
        raise ValueError(
            f"La hoja '{cruce_excel.HOJA_EXCEL}' no existe. Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[cruce_excel.HOJA_EXCEL]
    encabezados = [ws.cell(row=cruce_excel.FILA_ENCABEZADO, column=c).value for c in range(1, ws.max_column + 1)]
    col_rad = cruce_excel.encontrar_columna(encabezados, cruce_excel.COLUMNA_RADICADO)

    filas = []
    for fila in range(cruce_excel.FILA_ENCABEZADO + 1, ws.max_row + 1):
        radicado_crudo = ws.cell(row=fila, column=col_rad).value
        if radicado_crudo is None:
            continue
        radicado = re.sub(r"[\s\-]", "", str(radicado_crudo).strip())
        if not radicado.isdigit() or len(radicado) != 23:
            continue
        filas.append((fila, radicado))
    return filas


def procesar():
    if not cruce_excel.RUTA_EXCEL or not os.path.exists(cruce_excel.RUTA_EXCEL):
        logging.error("[Excel] No se encontro el archivo configurado en RUTA_EXCEL: %r", cruce_excel.RUTA_EXCEL)
        return

    filas_validas = leer_radicados_excel()
    logging.info("Excel: %d proceso(s) con radicado valido de 23 digitos.", len(filas_validas))

    radicados_exactos, radicados_base = _radicados_en_disco()
    logging.info(
        "Disco: %d carpeta(s) con radicado reconocible en %s.",
        len(radicados_exactos), cruce_excel.CARPETA_PROCESOS,
    )

    try:
        datos_extra = cruce_excel.leer_datos_faltantes_por_radicado()
    except Exception:
        datos_extra = {}

    encontrados = 0
    faltantes = []
    for fila, radicado in filas_validas:
        if radicado in radicados_exactos or radicado[:-1] in radicados_base:
            encontrados += 1
            continue
        faltantes.append((fila, radicado))

    logging.info(
        "Resultado: %d proceso(s) SI tienen carpeta en el disco (radicado igual o con distinto consecutivo), "
        "%d proceso(s) NO tienen ninguna carpeta.",
        encontrados, len(faltantes),
    )

    if not faltantes:
        logging.info("Todos los procesos del Excel tienen carpeta en el disco. No se genero reporte.")
        return

    faltantes.sort(key=lambda t: t[1])
    filas_reporte = []
    for fila, radicado in faltantes:
        datos = datos_extra.get(radicado, {})
        filas_reporte.append((
            fila,
            datos.get("cuenta", ""),
            radicado,
            datos.get("juzgado", ""),
            datos.get("demandado", ""),
            datos.get("estado", ""),
        ))

    if cruce_excel._escribir_csv_tolerante(
        ARCHIVO_REPORTE,
        ["Fila Excel", "Cuenta", "Radicado", "Juzgado", "Demandado", "Estado"],
        filas_reporte,
        "Faltan en el disco",
    ):
        logging.info("[Reporte] Guardado en: %s", ARCHIVO_REPORTE)

    logging.warning("%d proceso(s) del Excel NO tienen ninguna carpeta en el disco:", len(faltantes))
    for fila, radicado in faltantes:
        logging.warning("   - Fila %s del Excel: radicado %s", fila, radicado)


def main():
    configurar_logging()
    procesar()


if __name__ == "__main__":
    main()
